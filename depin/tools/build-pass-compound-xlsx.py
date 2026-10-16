import sys, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = sys.argv[1]
F = "Arial"
BLUE = Font(name=F, size=10, color="FF0000FF")
BLACK = Font(name=F, size=10)
GREEN = Font(name=F, size=10, color="FF008000")
BOLD = Font(name=F, size=10, bold=True)
TITLE = Font(name=F, size=13, bold=True)
NOTE = Font(name=F, size=9, italic=True, color="FF595959")
RED = Font(name=F, size=10, bold=True, color="FFC00000")
YEL = PatternFill("solid", fgColor="FFFFF2CC")
HEAD = PatternFill("solid", fgColor="FFDDEBF7")
thin = Side(style="thin", color="FFBFBFBF")
BOX = Border(bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")

NUM0 = '#,##0;(#,##0);"-"'
NUM1 = '#,##0.0;(#,##0.0);"-"'
NUM2 = '#,##0.00;(#,##0.00);"-"'
USD0 = '$#,##0;($#,##0);"-"'
USD2 = '$#,##0.00;($#,##0.00);"-"'
PCT = '0.0%;(0.0%);"-"'
MULT = '0.0"x";(0.0"x");"-"'

wb = Workbook()

def header(ws, row, labels):
    for i, t in enumerate(labels, 1):
        c = ws.cell(row=row, column=i, value=t)
        c.font = BOLD; c.fill = HEAD; c.alignment = Alignment(wrap_text=True, vertical="center")

# ---------------- 가정 ----------------
a = wb.active; a.title = "가정"
a["A1"] = "AI 이용권 재구매(복리) 시뮬레이션 — 가정"; a["A1"].font = TITLE
a["A2"] = "내부 검토용 · 대외 배포 · 고객 안내 금지. WLBN 가격은 가정(전망 아님). 지급 수량 · 가치는 사이트 약관상 비보장. 파란 글씨 · 노란 칸만 바꾼다."
a["A2"].font = NOTE
header(a, 3, ["항목", "값", "근거"])
rows = [
    (4, "정가 (RLUSD ≈ USD / 장)", 380, NUM0, True, "사이트 상품 페이지(정본 §4-1)"),
    (5, "첫 구매 할인율", 0.30, PCT, True, "서우 지정(10/16)"),
    (6, "첫 구매가 (USD / 장)", "=B4*(1-B5)", NUM2, False, "정가 × (1 − 할인율) = 266"),
    (7, "첫 구매 장수", 1, NUM0, True, "서우 지정(10/16)"),
    (8, "일 지급 (WLBN / 장)", 5, NUM0, True, "정본 §4-1"),
    (9, "지급 기간 (일)", 1095, NUM0, True, "정본 §4-1 — 발급일 기준 3년 · 이 시트의 계산 창도 같은 1,095일"),
    (10, "WLBN 가격 (USD / WLBN)", 1.0, USD2, True, "서우 지정 가정(10/16) — 전망 아님. 10/14 AMM 풀 화면 가격은 약 0.46"),
    (11, "재투자 켜기 (1 = 켬 · 0 = 끔)", 1, NUM0, True, "0 이면 「재구매 없이 보유」와 같아진다"),
    (12, "재구매 할인율", 0.30, PCT, True, "가정 — 재구매에도 같은 할인이 된다고 둠(할인 코드 조건 미확인)"),
    (13, "재구매 단가 (USD / 장)", "=B4*(1-B12)", NUM2, False, "정가 × (1 − 재구매 할인율)"),
    (14, "교환 수수료 (WLBN → RLUSD)", 0.005, PCT, True, "10/14 AMM 풀 수수료 0.5%(정본 §8). 이용권은 RLUSD 결제만이라 WLBN 을 바꿔야 한다"),
    (15, "슬리피지 (가정)", 0.0, PCT, True, "0 = 가격 영향 없음(낙관). 현실 점검 시트의 풀 깊이 계산 참고"),
    (16, "계정당 보유 상한 (장)", 100, NUM0, True, "정본 §4-1 — 계정당 최대 100장"),
    (17, "재구매 1장에 필요한 WLBN", "=B13/(B10*(1-B14)*(1-B15))", NUM2, False, "재구매 단가 ÷ (가격 × (1 − 수수료) × (1 − 슬리피지))"),
    (18, "클레임 보류 (일)", 1, NUM0, False, "정본 §4-1 — 적립분 24시간 보류. 일별 시트 수식에 1일로 들어 있다(값만 바꿔도 반영 안 됨)"),
    (19, "첫 구매일 (09:00 KST 스냅샷 뒤 구매 가정)", datetime.date(2026, 10, 16), 'yyyy-mm-dd', True, "날짜 계산용 — 09:00 전에 사면 하루 당겨진다"),
]
for r, lab, val, fmt, is_in, src in rows:
    a.cell(row=r, column=1, value=lab).font = BLACK
    c = a.cell(row=r, column=2, value=val); c.number_format = fmt
    if is_in:
        c.font = BLUE; c.fill = YEL
    else:
        c.font = BLACK
    a.cell(row=r, column=3, value=src).font = NOTE
a["A21"] = "계산 규칙"; a["A21"].font = BOLD
rules = [
    "① 첫 구매일 = 0일차. 매일 09:00 KST(00:00 UTC) 스냅샷에 보유 장수 × 5 WLBN 이 적립된다(첫 적립 = 1일차).",
    "② 적립분은 24시간 보류 → 다음 날부터 쓸 수 있다. 쓸 수 있는 잔고가 재구매 1장 값 이상이면 그날 바로 정수 장만큼 산다(상한 100장까지).",
    "③ 새로 산 장은 다음 날 스냅샷부터 적립된다. 3년 창 안에서는 어떤 장도 만료되지 않는다(첫 장이 1,095일차에 끝난다).",
    "④ 「평가액」은 손에 남은 WLBN(보류분 포함) × 가격이다. 재구매에 쓴 WLBN 과 이용권 자체의 남은 가치는 넣지 않는다.",
]
for i, t in enumerate(rules):
    c = a.cell(row=22 + i, column=1, value=t); c.font = BLACK; c.alignment = WRAP
    a.merge_cells(start_row=22 + i, start_column=1, end_row=22 + i, end_column=3)
    a.row_dimensions[22 + i].height = 28
a.column_dimensions["A"].width = 34; a.column_dimensions["B"].width = 14; a.column_dimensions["C"].width = 70

# ---------------- 일별 ----------------
d = wb.create_sheet("일별")
d["A1"] = "일별 계산 (1~1,095일차)"; d["A1"].font = TITLE
d["A2"] = "모든 칸은 수식. 0일차 = 첫 구매일."; d["A2"].font = NOTE
dh = ["일차", "월", "보유 이용권\n(그날 적립 대상)", "그날 적립\nWLBN", "쓸 수 있는 잔고\n(재구매 전)", "재구매 장수", "재구매에 쓴\nWLBN",
      "그날 끝\n미사용 잔고", "누적 적립\nWLBN", "평가액 USD\n(잔고+보류)", "재구매 없이\n누적 WLBN", "재구매 뒤\n보유 장수"]
header(d, 4, dh)
d.row_dimensions[4].height = 42
FIRST, LAST = 5, 5 + 1095 - 1
for i in range(1095):
    r = FIRST + i
    day = i + 1
    d.cell(row=r, column=1, value=day)
    d.cell(row=r, column=2, value=f"=MATCH(A{r},월별!$C$6:$C$41,1)")
    if day == 1:
        d.cell(row=r, column=3, value="=가정!$B$7")
        d.cell(row=r, column=5, value=0)
        d.cell(row=r, column=9, value=f"=D{r}")
    else:
        d.cell(row=r, column=3, value=f"=L{r-1}")
        d.cell(row=r, column=5, value=f"=H{r-1}+D{r-1}")
        d.cell(row=r, column=9, value=f"=I{r-1}+D{r}")
    d.cell(row=r, column=4, value=f"=C{r}*가정!$B$8")
    d.cell(row=r, column=6, value=f"=IF(가정!$B$11=1,MAX(0,MIN(가정!$B$16-C{r},INT(E{r}/가정!$B$17))),0)")
    d.cell(row=r, column=7, value=f"=F{r}*가정!$B$17")
    d.cell(row=r, column=8, value=f"=E{r}-G{r}")
    d.cell(row=r, column=10, value=f"=(H{r}+D{r})*가정!$B$10")
    d.cell(row=r, column=11, value=f"=A{r}*가정!$B$8*가정!$B$7")
    d.cell(row=r, column=12, value=f"=C{r}+F{r}")
    for col in range(1, 13):
        c = d.cell(row=r, column=col); c.font = BLACK
        c.number_format = {1: NUM0, 2: NUM0, 3: NUM0, 6: NUM0, 12: NUM0, 10: USD0}.get(col, NUM1)
for col, w in zip("ABCDEFGHIJKL", [7, 5, 12, 11, 13, 10, 12, 13, 13, 14, 13, 11]):
    d.column_dimensions[col].width = w
d.freeze_panes = "B5"

# ---------------- 월별 ----------------
m = wb.create_sheet("월별")
m["A1"] = "월별 표 (36개월 · 1,095일을 36등분)"; m["A1"].font = TITLE
m["A2"] = "WLBN 가격은 가정 시트 B10. 「평가액」= 손에 남은 WLBN × 가격(재구매에 쓴 몫 · 이용권 자체 가치 제외)."; m["A2"].font = NOTE
m["A3"] = "=\"가정: 첫 구매 \"&TEXT(가정!B7,\"0\")&\"장 · 구매가 $\"&TEXT(가정!B6,\"0\")&\" · WLBN $\"&TEXT(가정!B10,\"0.00\")&\" · 재투자 \"&IF(가정!B11=1,\"켬\",\"끔\")&\" · 상한 \"&TEXT(가정!B16,\"0\")&\"장\""
m["A3"].font = BOLD
mh = ["월", "끝 일차", "시작 일차", "월말 보유\n이용권(장)", "그 달 적립\nWLBN", "그 달\n재구매(장)", "누적 적립\nWLBN", "월말 손에 남은\nWLBN", "월말 평가액\nUSD",
      "재구매 없이\n누적 WLBN", "재구매 없이\n평가액 USD", "원금 대비\n(재구매)", "원금 대비\n(재구매 없이)"]
header(m, 5, mh); m.row_dimensions[5].height = 42
for k in range(36):
    r = 6 + k
    m.cell(row=r, column=1, value=k + 1)
    m.cell(row=r, column=2, value=f"=ROUND(A{r}*가정!$B$9/36,0)")
    m.cell(row=r, column=3, value=1 if k == 0 else f"=B{r-1}+1")
    m.cell(row=r, column=4, value=f"=INDEX(일별!$L${FIRST}:$L${LAST},B{r})")
    m.cell(row=r, column=5, value=f"=SUMIFS(일별!$D${FIRST}:$D${LAST},일별!$B${FIRST}:$B${LAST},A{r})")
    m.cell(row=r, column=6, value=f"=SUMIFS(일별!$F${FIRST}:$F${LAST},일별!$B${FIRST}:$B${LAST},A{r})")
    m.cell(row=r, column=7, value=f"=INDEX(일별!$I${FIRST}:$I${LAST},B{r})")
    m.cell(row=r, column=8, value=f"=INDEX(일별!$H${FIRST}:$H${LAST},B{r})+INDEX(일별!$D${FIRST}:$D${LAST},B{r})")
    m.cell(row=r, column=9, value=f"=H{r}*가정!$B$10")
    m.cell(row=r, column=10, value=f"=INDEX(일별!$K${FIRST}:$K${LAST},B{r})")
    m.cell(row=r, column=11, value=f"=J{r}*가정!$B$10")
    m.cell(row=r, column=12, value=f"=IF(가정!$B$6*가정!$B$7=0,0,I{r}/(가정!$B$6*가정!$B$7))")
    m.cell(row=r, column=13, value=f"=IF(가정!$B$6*가정!$B$7=0,0,K{r}/(가정!$B$6*가정!$B$7))")
    for col in range(1, 14):
        c = m.cell(row=r, column=col); c.font = BLACK
        c.number_format = {1: NUM0, 2: NUM0, 3: NUM0, 4: NUM0, 6: NUM0, 9: USD0, 11: USD0, 12: MULT, 13: MULT}.get(col, NUM0)
    if (k + 1) % 12 == 0:
        for col in range(1, 14):
            m.cell(row=r, column=col).border = Border(bottom=Side(style="medium", color="FF808080"))
for col, w in zip("ABCDEFGHIJKLM", [5, 8, 8, 11, 12, 9, 13, 14, 14, 12, 12, 11, 11]):
    m.column_dimensions[col].width = w
m.freeze_panes = "B6"

# ---------------- 요약 ----------------
s = wb.create_sheet("요약", 1)
s["A1"] = "요약 — 이용권 1장(30% 할인) · WLBN $1.0 가정 · 3년"; s["A1"].font = TITLE
s["A2"] = "내부 검토용 · 대외 배포 · 고객 안내 금지. 「복리」는 적립분을 팔아 이용권을 더 사는 경우의 계산일 뿐 약속 · 전망이 아니다."; s["A2"].font = NOTE
header(s, 4, ["항목", "재구매 없이 보유", "재구매(복리)", "메모"])
L = LAST
summ = [
    (5, "첫 구매 원금 (USD)", "=가정!B6*가정!B7", "=가정!B6*가정!B7", USD0, ""),
    (6, "3년 누적 적립 (WLBN)", f"=일별!K{L}", f"=일별!I{L}", NUM0, "재구매 없이 = 장당 5,475"),
    (7, "재구매에 쓴 WLBN", 0, f"=SUM(일별!G{FIRST}:G{L})", NUM0, "WLBN 을 팔아 RLUSD 로 산 몫"),
    (8, "재구매 장수 (합계)", 0, f"=SUM(일별!F{FIRST}:F{L})", NUM0, "첫 구매 제외"),
    (9, "3년 끝 보유 이용권 (장)", "=가정!B7", f"=일별!L{L}", NUM0, ""),
    (10, "3년 끝 손에 남은 WLBN", f"=일별!K{L}", f"=일별!H{L}+일별!D{L}", NUM0, "보류분(마지막 날 적립) 포함"),
    (11, "3년 끝 평가액 (USD)", "=B10*가정!B10", "=C10*가정!B10", USD0, "손에 남은 WLBN × 가정 가격"),
    (12, "원금 대비 배수", "=IF(B5=0,0,B11/B5)", "=IF(C5=0,0,C11/C5)", MULT, ""),
    (13, "상한(계정당 장수) 도달 일차", "-", f"=IFERROR(INDEX(일별!A{FIRST}:A{L},MATCH(가정!B16,일별!L{FIRST}:L{L},0)),\"미도달\")", NUM0, "도달 뒤엔 더 살 수 없어 복리가 멈춘다"),
    (14, "상한 도달 뒤 하루 적립 (WLBN)", "=가정!B7*가정!B8", "=가정!B16*가정!B8", NUM0, ""),
    (15, "3년 뒤에도 남는 적립 권리 (WLBN)", 0, f"=SUMPRODUCT(일별!F{FIRST}:F{L},일별!A{FIRST}:A{L})*가정!B8", NUM0, "재구매분은 산 날부터 1,095일이라 3년 창 밖으로 이어진다(평가액에 넣지 않음)"),
    (16, "재구매에 필요한 RLUSD (합계)", 0, "=C8*가정!B13", USD0, "현실 점검 시트의 풀 깊이와 비교"),
    (17, "첫 재구매 일차", "-", f"=IFERROR(INDEX(일별!A{FIRST}:A{L},MATCH(TRUE,INDEX(일별!F{FIRST}:F{L}>0,0),0)),\"없음\")", NUM0, "적립분이 재구매 1장 값에 닿고 24시간 보류가 풀린 날"),
    (18, "첫 재구매 날짜", "-", "=IF(ISNUMBER(C17),가정!B19+C17,\"-\")", 'yyyy-mm-dd', "가정 시트 첫 구매일 기준"),
]
for r, lab, b, c, fmt, memo in summ:
    s.cell(row=r, column=1, value=lab).font = BLACK
    for col, v in ((2, b), (3, c)):
        cell = s.cell(row=r, column=col, value=v); cell.font = BLACK; cell.number_format = fmt
    s.cell(row=r, column=4, value=memo).font = NOTE
for r in (11, 12, 17, 18):
    for col in (1, 2, 3):
        s.cell(row=r, column=col).font = BOLD

s["A20"] = "연도별"; s["A20"].font = BOLD
header(s, 21, ["연차", "재구매 없이 누적 WLBN", "재구매 누적 적립 WLBN", "재구매 연말 손에 남은 WLBN", "재구매 연말 보유 장수", "재구매 연말 평가액 USD", "재구매 없이 평가액 USD"])
for i, mrow in enumerate((17, 29, 41)):
    r = 22 + i
    s.cell(row=r, column=1, value=f"{i+1}년차").font = BLACK
    for col, src, fmt in ((2, f"=월별!J{mrow}", NUM0), (3, f"=월별!G{mrow}", NUM0), (4, f"=월별!H{mrow}", NUM0), (5, f"=월별!D{mrow}", NUM0), (6, f"=월별!I{mrow}", USD0), (7, f"=월별!K{mrow}", USD0)):
        c = s.cell(row=r, column=col, value=src); c.font = GREEN; c.number_format = fmt
for col, w in zip("ABCDEFG", [32, 18, 18, 22, 18, 20, 20]):
    s.column_dimensions[col].width = w
s.column_dimensions["D"].width = 48

# ---------------- 현실 점검 ----------------
z = wb.create_sheet("현실 점검")
z["A1"] = "현실 점검 — 이 계산이 실제로 성립하려면"; z["A1"].font = TITLE
z["A2"] = "계산은 「재구매가 언제든 같은 조건으로 된다」고 둔 것. 아래가 하나라도 막히면 복리는 그만큼 줄거나 멈춘다."; z["A2"].font = NOTE
header(z, 4, ["항목", "값", "설명"])
zr = [
    (5, "10/14 AMM 풀 RLUSD 잔고", 122000, USD0, True, "XRPL 탐색기 화면(정본 §8) — 가정 아닌 그날 값"),
    (6, "10/14 AMM 풀 WLBN 잔고", 264800, NUM0, True, "같은 화면"),
    (7, "풀 화면 가격 (RLUSD / WLBN)", "=B5/B6", USD2, False, "시트 가정($1.0)의 절반 아래"),
    (8, "재구매에 필요한 RLUSD (요약)", "=요약!C16", USD0, False, "요약 시트에서"),
    (9, "그 RLUSD 가 풀 RLUSD 에서 차지하는 비율", "=IF(B5=0,0,B8/B5)", PCT, False, "한 번에 빼면 가격이 크게 밀린다"),
    (10, "한 번에 바꾼다면 필요한 WLBN (상수곱 풀)", "=IF(B5<=B8,\"풀 부족\",B8*B6/(B5-B8)/(1-가정!B14))", NUM0, False, "y = R_r · x(1−f) / (R_w + x(1−f)) 를 x 로 푼 값"),
    (11, "그때 평균 교환가 (RLUSD / WLBN)", "=IF(ISNUMBER(B10),B8/B10,\"-\")", USD2, False, "화면 가격 대비 얼마나 손해인지"),
    (12, "평균 교환가 ÷ 화면 가격", "=IF(ISNUMBER(B11),B11/B7,\"-\")", PCT, False, "100% 보다 낮을수록 가격 영향이 크다"),
    (13, "발행 총량 (장)", 4000, NUM0, True, "정본 §4-1 — 완판 시 재구매 불가"),
    (14, "10/12 화면 판매 (장)", 101, NUM0, True, "정본 §4-1"),
    (15, "한 계정이 상한까지 사면 총량 대비", "=가정!B16/B13", PCT, False, "한 사람이 전체의 2.5%"),
]
for r, lab, v, fmt, is_in, memo in zr:
    z.cell(row=r, column=1, value=lab).font = BLACK
    c = z.cell(row=r, column=2, value=v); c.number_format = fmt
    if is_in:
        c.font = BLUE; c.fill = YEL
    else:
        c.font = BLACK
    z.cell(row=r, column=3, value=memo).font = NOTE
z["A17"] = "막히는 지점"; z["A17"].font = BOLD
blocks = [
    "① 교환: 이용권은 RLUSD 결제만(정본 §4-1) → 재구매하려면 WLBN 을 RLUSD 로 바꿔야 한다. 10/16 서우 확인 — 교환 정상 작동, 재구매는 보유자 선택. 남은 것 = 사이트 문구(G4 전 거래 금지)와 라이브 풀의 정리(정본 §8).",
    "② 출금 · 클레임: 적립분이 계정 포인트로만 쌓이는 단계면(FAQ 「출금 개시 전 계정 포인트」) 지갑으로 꺼내 팔 수 없다(정본 §4-1 확인 필요 ②).",
    "③ 가격: 시트는 WLBN $1.0 을 고정으로 둔다. 재구매가 늘수록 같은 사람이 계속 파는 쪽이라 가격을 누른다 — 고정 가격은 낙관 쪽 가정이다.",
    "④ 할인: 첫 구매 30% 할인이 재구매에도 된다고 뒀다. 할인 코드 조건이 1회 · 1인이면 재구매 단가는 380 → 필요한 WLBN 이 약 1.4배.",
    "⑤ 수량: 발행 4,000장 · 계정당 100장. 매진되면 재구매 불가, 상한 도달 뒤엔 복리가 멈추고 하루 500 WLBN 고정 적립만 남는다.",
    "⑥ 용도: 이용권은 AI 팩토리 콘텐츠 결제용이고 결제에 쓴 토큰은 소각된다. 「적립분 팔아 이용권 더 사기」는 상품 설계 의도 밖의 사용이다.",
    "⑦ 세금 · 규제: 적립 시점 소득 과세 · 투자 권유 해당 여부는 세무 · 카운슬 확인 전. 「복리 · 수익률 · 월 ○달러」 표현은 대외 금지(playbook).",
]
for i, t in enumerate(blocks):
    c = z.cell(row=18 + i, column=1, value=t); c.font = BLACK; c.alignment = WRAP
    z.merge_cells(start_row=18 + i, start_column=1, end_row=18 + i, end_column=3)
    z.row_dimensions[18 + i].height = 30
z.column_dimensions["A"].width = 40; z.column_dimensions["B"].width = 14; z.column_dimensions["C"].width = 64

wb.save(OUT)
print("saved", OUT)
