import argparse, re
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ap = argparse.ArgumentParser()
ap.add_argument("src"); ap.add_argument("out")
ap.add_argument("--title", default="업종별 측정기 영업 핵심 메시지 — 30개 업종 표")
ap.add_argument("--seo", action="store_true",
                help="서우 정리본 형식: 머리글 · 부제 없음, 업종 번호, 마지막 열 「비고」 빈칸(md 의 주의 열은 싣지 않음)")
args = ap.parse_args()
SRC, OUT = args.src, args.out
FONT = "Malgun Gothic"
INK, ACCENT, MUTE = "1B1B48", "4D4DCE", "6B6B8C"
HEAD_FILL, GROUP_FILL, ZEBRA, RULE = "E8E8F4", "D9D9F0", "F8F8FC", "C8C8DA"

# --- parse the industry tables from the md: "## 그룹 (n)" + 5-column rows ---
groups, cur = [], None
for line in open(SRC, encoding="utf-8").read().splitlines():
    m = re.match(r"^## (.+\((\d+)\))\s*$", line)
    if m:
        cur = (m.group(1), [], int(m.group(2))); groups.append(cur); continue
    if line.startswith("## "):
        cur = None; continue
    if cur and line.startswith("|") and not line.startswith("|---") and not line.startswith("| 업종"):
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if len(cells) == 5:
            cur[1].append(cells)
assert groups and all(len(g[1]) == g[2] for g in groups), [(g[0], len(g[1])) for g in groups]
print(len(groups), "groups,", sum(len(g[1]) for g in groups), "rows")

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), fill)
    tcPr.append(sh)

def borders(table):
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:color"), RULE)
        b.append(e)
    tblPr.append(b)

def write(cell, text, size=9, bold=False, color=INK):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(0)
    r = p.add_run(text); r.font.size = Pt(size); r.bold = bold; r.font.color.rgb = RGBColor.from_string(color)
    r.font.name = FONT; r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def set_width(cell, cm):
    cell.width = Cm(cm)

doc = Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
for side in ("left_margin", "right_margin"): setattr(sec, side, Cm(1.5))
sec.top_margin, sec.bottom_margin = (Cm(1.2), Cm(1.1)) if args.seo else (Cm(1.4), Cm(1.4))
st = doc.styles["Normal"]; st.font.name = FONT; st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT); st.font.size = Pt(9)

# header line (서우 형식은 머리글 · 부제 없음)
if not args.seo:
  hp = sec.header.paragraphs[0]
  hr = hp.add_run("wellbian 내부 · 업종별 측정기 영업 핵심 메시지 · 1:1 영업용 · 게시 · 공표 · 전달 금지")
  hr.font.size = Pt(8); hr.font.color.rgb = RGBColor.from_string(MUTE); hr.font.name = FONT; hr._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

t = doc.add_paragraph(); tr = t.add_run(args.title)
tr.bold = not args.seo; tr.font.size = Pt(16); tr.font.color.rgb = RGBColor.from_string(INK); tr.font.name = FONT; tr._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
t.paragraph_format.space_after = Pt(2)
if not args.seo:
  s = doc.add_paragraph(); sr = s.add_run("측정 항목: 초미세먼지 · 미세먼지 · CO₂ · TVOC · 온도 · 습도(판매 페이지 사양). 건강 효과 · 절감액 · 수익 · 법정 측정 대체는 말하지 않는다. 2026-10-19")
  sr.font.size = Pt(8.5); sr.font.color.rgb = RGBColor.from_string(MUTE); sr.font.name = FONT; sr._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
  s.paragraph_format.space_after = Pt(6)
else:
  t.paragraph_format.space_after = Pt(6)

W = [3.6, 4.6, 8.2, 3.4, 6.9]   # cm, sum 26.7 = page 29.7 - 3.0
heads = ["업종", "현장의 불편", "핵심 메시지", "어디에", "주의"]
if args.seo:   # 서우 정리본 PDF 의 열 비율 — 빈칸인 비고를 줄여 핵심 메시지 한 글자 줄넘김을 없앰
    W = [4.0, 5.0, 9.4, 4.5, 3.8]
    heads[4] = "비고"
tbl = doc.add_table(rows=1, cols=5)
tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
tbl.autofit = False
borders(tbl)
tblPr = tbl._tbl.tblPr
lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
tw = OxmlElement("w:tblW"); tw.set(qn("w:w"), str(int(sum(W) * 567))); tw.set(qn("w:type"), "dxa")
for old_tw in tblPr.findall(qn("w:tblW")): tblPr.remove(old_tw)
tblPr.append(tw)
grid = tbl._tbl.tblGrid
for gc, w in zip(grid.findall(qn("w:gridCol")), W): gc.set(qn("w:w"), str(int(w * 567)))
for i, col in enumerate(tbl.columns): col.width = Cm(W[i])
for i, h in enumerate(heads):
    c = tbl.rows[0].cells[i]; write(c, h, size=9, bold=True); shade(c, HEAD_FILL); set_width(c, W[i])
# repeat header row on each page
trPr = tbl.rows[0]._tr.get_or_add_trPr(); th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trPr.append(th)

zebra, n = False, 0
for title, rows, _ in groups:
    gr = tbl.add_row()
    merged = gr.cells[0].merge(gr.cells[4])
    write(merged, title, size=9.5, bold=True, color=ACCENT); shade(merged, GROUP_FILL)
    gtrPr = gr._tr.get_or_add_trPr(); cs = OxmlElement("w:cantSplit"); gtrPr.append(cs)
    for cells in rows:
        n += 1
        if args.seo: cells = [f"{n}. {cells[0]}"] + cells[1:4] + [""]
        r = tbl.add_row()
        rp = r._tr.get_or_add_trPr(); cs2 = OxmlElement("w:cantSplit"); rp.append(cs2)
        for i, txt in enumerate(cells):
            c = r.cells[i]; write(c, txt, size=9, bold=(i == 0)); set_width(c, W[i])
            if zebra: shade(c, ZEBRA)
        zebra = not zebra
doc.save(OUT)
print("saved", OUT)
