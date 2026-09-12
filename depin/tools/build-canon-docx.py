#!/usr/bin/env python3
"""site-canon md → DOCX (사이트 개발자·법무 전달용).

정본은 depin/content/site-canon-0910.md 하나뿐이다. 이 스크립트는 그 md 를 읽어 docx 를 만들 뿐,
문장을 따로 갖지 않는다 — md 를 고친 뒤 다시 돌리면 docx 가 따라온다(사본이 갈라지지 않게).

  python3 depin/tools/build-canon-docx.py [src.md] [out.docx] [--portrait]

가로(기본)는 표가 넓은 정본 스냅샷용, --portrait 는 본문 위주 검토 문서용이다.
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ARGS = [a for a in sys.argv[1:] if not a.startswith("--")]
PORTRAIT = "--portrait" in sys.argv[1:]
SRC = Path(ARGS[0] if len(ARGS) > 0 else "depin/content/site-canon-0910.md")
OUT = Path(ARGS[1] if len(ARGS) > 1 else "depin/content/site-canon-0910.docx")

KO_FONT = "Malgun Gothic"
INK = RGBColor(0x1B, 0x1B, 0x48)
MUTE = RGBColor(0x6B, 0x6B, 0x8C)
HEAD_FILL = "E8E8F4"

INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")


def set_font(run, size=None, bold=None, color=None, mono=False):
    run.font.name = "Consolas" if mono else KO_FONT
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), KO_FONT)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_inline(par, text, size, color=INK):
    """**굵게** 와 `코드` 만 해석한다. 나머지 마크다운은 글자 그대로."""
    for tok in INLINE.split(text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            set_font(par.add_run(tok[2:-2]), size, bold=True, color=color)
        elif tok.startswith("`") and tok.endswith("`"):
            set_font(par.add_run(tok[1:-1]), size, color=color, mono=True)
        else:
            set_font(par.add_run(tok), size, color=color)


def shade(cell, fill):
    tcpr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def split_row(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def add_table(doc, rows):
    header, body = rows[0], [r for r in rows[1:] if not re.fullmatch(r"[\s|:-]+", "|".join(r))]
    ncol = len(header)
    t = doc.add_table(rows=1 + len(body), cols=ncol)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(header):
        c = t.rows[0].cells[j]
        c.text = ""
        add_inline(c.paragraphs[0], h, 9)
        for r in c.paragraphs[0].runs:
            r.font.bold = True
        shade(c, HEAD_FILL)
    for i, row in enumerate(body, start=1):
        for j in range(ncol):
            c = t.rows[i].cells[j]
            c.text = ""
            add_inline(c.paragraphs[0], row[j] if j < len(row) else "", 9)
    doc.add_paragraph()


def build(md_text):
    doc = Document()
    sec = doc.sections[0]
    if PORTRAIT:
        sec.orientation = WD_ORIENT.PORTRAIT
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    else:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(1.8))
    normal = doc.styles["Normal"]
    normal.font.name = KO_FONT
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), KO_FONT)
    normal.font.size = Pt(10)

    lines = md_text.splitlines()
    i, table, para, bullet = 0, [], [], []

    def flush_para():
        nonlocal para
        if para:
            p = doc.add_paragraph()
            add_inline(p, " ".join(para), 10)
            para = []

    def flush_bullet():
        """불릿이 여러 줄로 접혀 있어도 한 항목으로 합친다 —
        합치지 않으면 줄 경계에서 **강조** 가 갈라져 별표가 그대로 찍힌다."""
        nonlocal bullet
        if bullet:
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, " ".join(bullet), 10)
            bullet = []

    def flush_table():
        nonlocal table
        if table:
            add_table(doc, table)
            table = []

    while i < len(lines):
        ln = lines[i]
        if ln.startswith("|"):
            flush_bullet()
            flush_para()
            table.append(split_row(ln))
        else:
            flush_table()
            if ln.startswith("# "):
                flush_bullet()
                flush_para()
                p = doc.add_paragraph()
                set_font(p.add_run(ln[2:].strip()), 18, bold=True, color=INK)
                p = doc.add_paragraph()
                set_font(p.add_run("생성 문서 — 정본은 md 파일입니다. 이 문서를 고치지 말고 md 를 고친 뒤 다시 생성하세요 "
                                   "(depin/tools/build-canon-docx.py)."), 8.5, color=MUTE)
            elif ln.startswith("## "):
                flush_bullet()
                flush_para()
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(14)
                set_font(p.add_run(ln[3:].strip()), 13, bold=True, color=INK)
            elif ln.startswith("- "):
                flush_bullet()
                flush_para()
                bullet.append(ln[2:].strip())
            elif not ln.strip():
                flush_bullet()
                flush_para()
            elif bullet:
                bullet.append(ln.strip())
            else:
                para.append(ln.strip())
        i += 1
    flush_bullet()
    flush_para()
    flush_table()
    return doc


if __name__ == "__main__":
    doc = build(SRC.read_text(encoding="utf-8"))
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"wrote {OUT}")
