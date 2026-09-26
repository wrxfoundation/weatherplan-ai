#!/usr/bin/env python3
"""조사 보고서 md → DOCX (서우 전달용, 2026-09-26 「Docx화해줘」 → 「위계질서에 맞춰 디자인 / 폰트 정리」).

정본은 md 다. 이 스크립트는 md 를 읽어 docx 를 만들 뿐 문장을 따로 갖지 않는다 — md 를 고친 뒤 다시 돌리면
docx 가 따라온다(build-canon-docx.py 와 같은 원칙 · 같은 서체와 색).

  python3 depin/tools/build-report-docx.py "reports/XRPL 인플루언서와 DePIN 고객 확보.md" [out.docx] [--no-number]

out 을 빼면 md 옆에 같은 이름의 .docx 를 만든다.

위계(한 서체 · 맑은 고딕, 크기 · 굵기 · 색으로만 나눈다)
  표지 라벨 9 굵게(남색 강조) → 문서 제목 24 굵게(파일 이름) → 요지 13.5(md 의 # 문장) → 요약 · 목차
  절 ## 15 굵게, 번호 1 · 2 …(강조색) → 소절 ### 11.5 굵게, 번호 7.1 … → 문단 첫 굵은 구절 10 굵게 → 본문 10
  → 근거 표시([1차 원문] · [검색요약] · [내부] · 「원출처 확인 필요」 · unverified (aggregator) 등) 8.5 회색
  「[외부 근거]」 「[추론]」 은 작은 꼬리표로, [추론] 문단은 옅은 바탕으로 외부 근거 문단과 구별한다.
  링크는 밑줄 없이 강조색(출처가 본문보다 앞서 보이지 않게). 표는 머리글 반복 · 첫 열 굵게(4열 이상) · 열 너비 자동.
  열 6개 이상인 표가 든 ## 절은 가로 쪽. 첫 쪽은 표지 · 요약 · 목차(누르면 그 절로), 본문은 둘째 쪽부터.
  줄 높이는 고정값(본문 10/16pt) — Word 의 맑은 고딕과 대체 서체에서 쪽 배치가 달라지지 않게.
"""
import re
import sys
import unicodedata
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Pt, RGBColor

KO_FONT = "Malgun Gothic"
INK = "1B1B48"         # 본문 · 제목
ACCENT = "4D4DCE"      # 브랜드 남색 — 번호 · 라벨 · 링크 · 근거 꼬리표
MUTE = "6B6B8C"        # 근거 표시 · 머리글 · 쪽 번호
RULE = "C8C8DA"        # 표 선
HEAD_FILL = "E8E8F4"   # 표 머리글
ZEBRA = "F8F8FC"       # 표 짝수 줄
EVID_CHIP = "ECECFA"   # [외부 근거] 꼬리표 바탕
INFER_TXT = "8A5A00"   # [추론] 꼬리표 글자
INFER_CHIP = "F6EBD2"  # [추론] 꼬리표 바탕
INFER_BOX = "FBF7EF"   # [추론] 문단 바탕
MARGIN = Cm(1.8)

# 크기(pt) — 한곳에서 위계를 정한다
S_TITLE, S_SUB, S_H2, S_H3, S_BODY, S_LEAD, S_TAG, S_META = 24, 13.5, 15, 11.5, 10, 10.5, 8.5, 9
# 줄 높이(pt, 고정) — 서체가 바뀌어도(맑은 고딕 · Mac 대체 서체) 쪽 배치가 같게
# 고정 줄 높이는 글자 높이(맑은 고딕 ≈ 1.34em)보다 커야 위아래가 잘리지 않는다 — 제목 24pt 는 34pt
L_TITLE, L_SUB, L_H2, L_H3, L_BODY, L_LEAD, L_CHECK = 34, 20, 22, 17, 16, 16.5, 15.5

REPO = Path(__file__).resolve().parents[2]
ARGS = [a for a in sys.argv[1:] if not a.startswith("--")]
NUMBER = "--no-number" not in sys.argv[1:]
if not ARGS:
    sys.exit(__doc__)
SRC = Path(ARGS[0])
OUT = Path(ARGS[1]) if len(ARGS) > 1 else SRC.with_suffix(".docx")

TAGS = r"\[(?:1차 원문|검색요약|내부|일화|벤더 주장|회사 주장|업체 주장|분석가 의견)\]|「원출처 확인 필요」|unverified \(aggregator\)|stale 가능"
TOKEN = re.compile(
    r"\[(?P<lt>[^\]]+)\]\((?P<url>[^)\s]+)\)"
    r"|(?P<chip>\[(?:외부 근거|추론)\])"
    rf"|(?P<tag>{TAGS})"
    r"|\*\*(?P<b>.+?)\*\*"
    r"|`(?P<c>[^`]+)`"
    r"|(?<![*\w])\*(?P<i>[^\s*][^*]*?)\*(?![*\w])"
)


# ── 글자 ─────────────────────────────────────────────────────────────
def rpr(size, bold=False, italic=False, color=INK, underline=False, fill=None):
    """w:rPr — 스키마 순서(rFonts · b · i · color · sz · u · shd). 모든 글자는 한 서체."""
    el = OxmlElement("w:rPr")
    f = OxmlElement("w:rFonts")
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        f.set(qn(a), KO_FONT)
    el.append(f)
    if bold:
        el.append(OxmlElement("w:b"))
        el.append(OxmlElement("w:bCs"))
    if italic:
        el.append(OxmlElement("w:i"))
        el.append(OxmlElement("w:iCs"))
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    el.append(c)
    for tag in ("w:sz", "w:szCs"):
        s = OxmlElement(tag)
        s.set(qn("w:val"), str(int(round(size * 2))))
        el.append(s)
    if underline:
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        el.append(u)
    if fill:
        sh = OxmlElement("w:shd")
        sh.set(qn("w:val"), "clear")
        sh.set(qn("w:color"), "auto")
        sh.set(qn("w:fill"), fill)
        el.append(sh)
    return el


def run_el(text, size, **kw):
    r = OxmlElement("w:r")
    r.append(rpr(size, **kw))
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    r.append(t)
    return r


def add_link(p, url, text, size, bold=False, italic=False):
    rid = p.part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rid)
    h.set(qn("w:history"), "1")
    m = re.fullmatch(r"\*\*(.+)\*\*", text)
    if m:
        text, bold = m.group(1), True
    h.append(run_el(text, size, bold=bold, italic=italic, color=ACCENT))
    p._p.append(h)


def chip(p, label, size):
    """[외부 근거] · [추론] — 본문보다 작은 꼬리표(바탕색 안에 좁은 여백)."""
    infer = label == "추론"
    p._p.append(run_el(f"\u2009{label}\u2009", max(size - 2, 7.5), bold=True,
                       color=INFER_TXT if infer else ACCENT, fill=INFER_CHIP if infer else EVID_CHIP))


def inline(p, text, size, bold=False, italic=False, color=INK):
    """링크 · 꼬리표 · 근거 표시 · **굵게** · *기울임* · `코드`. 굵게 안도 다시 읽는다."""
    pos = 0
    for m in TOKEN.finditer(text):
        if m.start() > pos:
            p._p.append(run_el(text[pos:m.start()], size, bold=bold, italic=italic, color=color))
        if m.group("url"):
            add_link(p, m.group("url"), m.group("lt"), size, bold, italic)
        elif m.group("chip"):
            chip(p, m.group("chip")[1:-1], size)
        elif m.group("tag"):
            p._p.append(run_el(m.group("tag"), max(size - 1.5, 7), italic=italic, color=MUTE))
        elif m.group("b") is not None:
            inline(p, m.group("b"), size, bold=True, italic=italic, color=color)
        elif m.group("c") is not None:
            p._p.append(run_el(m.group("c"), size, bold=bold, italic=italic, color=MUTE))
        else:
            inline(p, m.group("i"), size, bold=bold, italic=True, color=color)
        pos = m.end()
    if pos < len(text):
        p._p.append(run_el(text[pos:], size, bold=bold, italic=italic, color=color))


def plain(text):
    """표시되는 글자만 — 폭 계산 · 목차용."""
    text = re.sub(r"\[([^\]]+)\]\([^)\s]+\)", r"\1", text)
    return text.replace("**", "").replace("`", "")


def units(text):
    """반각 1 · 전각(한글 등) 2."""
    return sum(2 if unicodedata.east_asian_width(ch) in "WF" else 1 for ch in text)


# ── 문단 꾸밈 ──────────────────────────────────────────────────────────
PPR_AFTER_SHD = ("w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct",
                 "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd",
                 "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents",
                 "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment", "w:textboxTightWrap",
                 "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange")


def tint(p, fill):
    """옅은 바탕 + 같은 색 테두리로 안쪽 여백(옆 줄무늬가 아니라 상자 전체)."""
    ppr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    for side in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:space"), "6")
        e.set(qn("w:color"), fill)
        bdr.append(e)
    ppr.insert_element_before(bdr, "w:shd", *PPR_AFTER_SHD)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    ppr.insert_element_before(shd, *PPR_AFTER_SHD)


def bookmark(p, name, bid):
    s = OxmlElement("w:bookmarkStart")
    s.set(qn("w:id"), str(bid))
    s.set(qn("w:name"), name)
    e = OxmlElement("w:bookmarkEnd")
    e.set(qn("w:id"), str(bid))
    ppr = p._p.find(qn("w:pPr"))
    if ppr is not None:
        ppr.addnext(s)
    else:
        p._p.insert(0, s)
    p._p.append(e)


def field(p, instr, size, color):
    def fld(kind):
        r = OxmlElement("w:r")
        r.append(rpr(size, color=color))
        c = OxmlElement("w:fldChar")
        c.set(qn("w:fldCharType"), kind)
        r.append(c)
        return r

    p._p.append(fld("begin"))
    r = OxmlElement("w:r")
    r.append(rpr(size, color=color))
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instr} "
    r.append(it)
    p._p.append(r)
    p._p.append(fld("separate"))
    p._p.append(run_el("1", size, color=color))
    p._p.append(fld("end"))


def hang(p, width):
    """번호 뒤 줄바꿈이 번호 오른쪽에 맞춰지게."""
    pf = p.paragraph_format
    pf.left_indent = (pf.left_indent or 0) + width
    pf.first_line_indent = -width
    pf.tab_stops.add_tab_stop(pf.left_indent)


# ── 표 ───────────────────────────────────────────────────────────────
TCPR_AFTER_SHD = ("w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark",
                  "w:headers", "w:cellIns", "w:cellDel", "w:cellMerge", "w:tcPrChange")


def shade(cell, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    cell._element.get_or_add_tcPr().insert_element_before(shd, *TCPR_AFTER_SHD)


def col_widths(rows, total, size):
    """열 너비(dxa) — 본문 평균 글자 폭에 비례, 머리글 · 끊기지 않는 반각 덩어리(핸들 등)가 한 줄에 들어갈 최소 폭은 지킨다."""
    n = len(rows[0])
    unit = size * 10  # 반각 한 글자 ≈ 0.5em = size*10 dxa
    pad = 2 * 90 + 60
    head = [units(plain(h)) for h in rows[0]]
    body = rows[1:] or rows[:1]
    raw, minw = [], []
    for j in range(n):
        cells = [plain(r[j] if j < len(r) else "") for r in body]
        lens = [units(c) for c in cells]
        tok = max((len(t) for c in cells for t in re.findall(r"[!-~]+", c)), default=0)
        raw.append(max(sum(lens) / len(lens), head[j], 4))
        minw.append(max((max(head[j], 4) + 1) * unit, min(tok, 18) * size * 12 + 60) + pad)  # 반각 실폭 ≈ 0.6em
    fixed = set()
    while True:
        rem = total - sum(minw[j] for j in fixed)
        free = [j for j in range(n) if j not in fixed]
        s = sum(raw[j] for j in free)
        w = {j: minw[j] for j in fixed}
        w.update({j: rem * raw[j] / s for j in free})
        low = [j for j in free if w[j] < minw[j]]
        if not low:
            break
        fixed.update(low)
    out = [int(w[j]) for j in range(n)]
    out[-1] += total - sum(out)
    return out


def add_table(doc, rows, total):
    n = len(rows[0])
    size = 8.5 if n <= 5 else 8
    widths = col_widths(rows, total, size)
    t = doc.add_table(rows=len(rows), cols=n)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    tpr = t._tbl.tblPr
    tw = tpr.find(qn("w:tblW"))
    if tw is None:
        tw = OxmlElement("w:tblW")
        tpr.insert_element_before(tw, "w:jc", "w:tblCellSpacing", "w:tblInd", "w:tblBorders", "w:shd",
                                  "w:tblLayout", "w:tblCellMar", "w:tblLook")
    tw.set(qn("w:type"), "dxa")
    tw.set(qn("w:w"), str(total))
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), RULE)
        borders.append(e)
    tpr.insert_element_before(borders, "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook",
                              "w:tblCaption", "w:tblDescription")
    mar = OxmlElement("w:tblCellMar")
    for side, v in (("top", 60), ("left", 90), ("bottom", 60), ("right", 90)):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:w"), str(v))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    tpr.insert_element_before(mar, "w:tblLook", "w:tblCaption", "w:tblDescription")
    for j, gc in enumerate(t._tbl.tblGrid.findall(qn("w:gridCol"))):
        gc.set(qn("w:w"), str(widths[j]))
    for i, row in enumerate(t.rows):
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:cantSplit"))
        if i == 0:
            trpr.append(OxmlElement("w:tblHeader"))
        cells = rows[i] + [""] * (n - len(rows[i]))
        for j, cell in enumerate(row.cells):
            cell.width = Emu(widths[j] * 635)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = Pt(12 if size >= 8.5 else 11.5)
            inline(p, cells[j], size, bold=(i == 0 or (j == 0 and n >= 4)))
            if i == 0:
                shade(cell, HEAD_FILL)
            elif i % 2 == 0:
                shade(cell, ZEBRA)
    gap = doc.add_paragraph()
    gap.paragraph_format.space_after = Pt(6)


# ── md 읽기 ──────────────────────────────────────────────────────────
def split_row(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def parse(md):
    """블록 목록. ## · ### 에는 책갈피 이름과 번호(1 · 7.1 …)를 붙인다."""
    blocks, para, k, n2, n3 = [], [], 0, 0, 0
    lines = md.splitlines()

    def flush():
        if para:
            blocks.append(("p", " ".join(s.strip() for s in para)))
            para.clear()

    i = 0
    while i < len(lines):
        ln = lines[i]
        if ln.startswith("|"):
            flush()
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                rows.append(split_row(lines[i]))
                i += 1
            blocks.append(("table", [r for r in rows if not re.fullmatch(r"[\s|:-]+", "|".join(r))]))
            continue
        m = re.match(r"(#{1,3}) (.+)", ln)
        if m:
            flush()
            level = len(m.group(1))
            if level == 1:
                blocks.append(("h1", m.group(2).strip(), "", ""))
            else:
                k += 1
                if level == 2:
                    n2, n3 = n2 + 1, 0
                    num = f"{n2}"
                else:
                    n3 += 1
                    num = f"{n2}.{n3}"
                blocks.append((f"h{level}", m.group(2).strip(), f"s{k}", num if NUMBER else ""))
        elif re.match(r"- \[[ xX]\] ", ln):
            flush()
            blocks.append(("check", ln[6:].strip()))
        elif re.match(r"[-*] ", ln):
            flush()
            blocks.append(("bullet", ln[2:].strip()))
        elif not ln.strip():
            flush()
        else:
            para.append(ln)
        i += 1
    flush()
    return blocks


# ── 문서 ─────────────────────────────────────────────────────────────
def no_theme(rfonts):
    for key in list(rfonts.attrib):
        if "heme" in key:
            del rfonts.attrib[key]
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(a), KO_FONT)


def setup(doc, title):
    d = doc.styles.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr"))
    rf = d.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        d.insert(0, rf)
    no_theme(rf)
    lang = d.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        d.append(lang)
    lang.set(qn("w:val"), "ko-KR")
    lang.set(qn("w:eastAsia"), "ko-KR")
    normal = doc.styles["Normal"]
    normal.font.size = Pt(S_BODY)
    normal.font.color.rgb = RGBColor.from_string(INK)
    no_theme(normal.element.get_or_add_rPr().get_or_add_rFonts())
    normal.paragraph_format.line_spacing = Pt(L_BODY)
    for name, size, before, after, pitch in (("Heading 1", S_TITLE, 0, 4, L_TITLE), ("Heading 2", S_H2, 22, 8, L_H2),
                                             ("Heading 3", S_H3, 14, 5, L_H3)):
        st = doc.styles[name]
        no_theme(st.element.get_or_add_rPr().get_or_add_rFonts())
        st.font.size, st.font.bold, st.font.italic = Pt(size), True, False
        st.font.color.rgb = RGBColor.from_string(INK)
        pf = st.paragraph_format
        pf.space_before, pf.space_after, pf.keep_with_next = Pt(before), Pt(after), True
        pf.line_spacing = Pt(pitch)
    zoom = doc.settings.element.find(qn("w:zoom"))
    if zoom is not None and zoom.get(qn("w:percent")) is None:
        zoom.set(qn("w:percent"), "100")  # 기본 틀의 zoom 에 필수 속성이 빠져 있다(스키마 검사)
    doc.core_properties.title = title
    doc.core_properties.author = "wellbian"
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, MARGIN)
    sec.different_first_page_header_footer = True  # 표지에는 머리글 · 쪽 번호 없음
    hp = sec.header.paragraphs[0]
    hp.paragraph_format.space_after = Pt(0)
    hp._p.append(run_el(f"wellbian 내부 · {title} · {date.today().isoformat()} · 수치·명단은 원출처 확인 전 대외 인용 금지",
                        7.5, color=MUTE))
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field(fp, "PAGE", 8, MUTE)
    fp._p.append(run_el(" / ", 8, color=MUTE))
    field(fp, "NUMPAGES", 8, MUTE)


def body_width(doc):
    s = doc.sections[-1]
    return int((s.page_width - s.left_margin - s.right_margin) / 635)


def orient(doc, landscape):
    s = doc.sections[-1]
    if (s.orientation == WD_ORIENT.LANDSCAPE) == landscape:
        return
    s = doc.add_section(WD_SECTION.NEW_PAGE)
    s.different_first_page_header_footer = False
    s.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    s.page_width, s.page_height = (Cm(29.7), Cm(21.0)) if landscape else (Cm(21.0), Cm(29.7))
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(s, side, MARGIN)


def para(doc, text, size=S_BODY, after=7, line=L_BODY):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = Pt(line)
    inline(p, text, size)
    return p


def label(doc, text, before=0, after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.keep_with_next = True
    p._p.append(run_el(text, S_META, bold=True, color=ACCENT))
    return p


def heading(doc, b, bid):
    kind, text, name, num = b
    level = 2 if kind == "h2" else 3
    size = S_H2 if level == 2 else S_H3
    h = doc.add_paragraph(style=f"Heading {level}")
    if num:
        hang(h, Cm(1.0 if level == 2 else 1.2))
        h._p.append(run_el(f"{num}\t", size, bold=True, color=ACCENT))
    inline(h, text, size, bold=True)
    bookmark(h, name, bid)
    return h


def contents(doc, heads):
    label(doc, "목차", before=12, after=5)
    for kind, text, name, num in heads:
        e = doc.add_paragraph()
        e.paragraph_format.space_after = Pt(1.5 if kind == "h2" else 1)
        e.paragraph_format.line_spacing = Pt(13.5 if kind == "h2" else 12)
        e.paragraph_format.keep_with_next = (kind, name) != (heads[-1][0], heads[-1][2])  # 목차는 한 덩어리로
        size = 9.5 if kind == "h2" else 8.5
        if kind == "h3":
            e.paragraph_format.left_indent = Cm(0.8)
        if num:
            hang(e, Cm(0.8 if kind == "h2" else 0.9))
        h = OxmlElement("w:hyperlink")
        h.set(qn("w:anchor"), name)
        h.set(qn("w:history"), "1")
        if num:
            h.append(run_el(f"{num}\t", size, bold=kind == "h2", color=ACCENT))
        h.append(run_el(plain(text), size, color=INK if kind == "h2" else MUTE))
        e._p.append(h)


def build(md, title, src_label):
    doc = Document()
    setup(doc, title)
    blocks = parse(md)

    # ## 절 단위로 묶어 가로 쪽 여부를 정한다
    groups, cur = [], []
    for b in blocks:
        if b[0] == "h2" and cur:
            groups.append(cur)
            cur = []
        cur.append(b)
    groups.append(cur)
    heads = [b for b in blocks if b[0] in ("h2", "h3")]

    bid, lead_done, first_h2 = 0, False, True
    for g in groups:
        orient(doc, any(b[0] == "table" and len(b[1][0]) >= 6 for b in g))
        for b in g:
            kind = b[0]
            if kind == "h1":
                k = doc.add_paragraph()
                k.paragraph_format.space_before = Pt(4)
                k.paragraph_format.space_after = Pt(10)
                k._p.append(run_el("조사 보고서", S_META, bold=True, color=ACCENT))
                k._p.append(run_el(f"   wellbian 내부 · {date.today().isoformat()}", S_META, color=MUTE))
                t = doc.add_paragraph(style="Heading 1")
                t._p.append(run_el(title, S_TITLE, bold=True))
                s = doc.add_paragraph()
                s.paragraph_format.space_after = Pt(16)
                s.paragraph_format.line_spacing = Pt(L_SUB)
                inline(s, b[1], S_SUB)
            elif kind in ("h2", "h3"):
                bid += 1
                h = heading(doc, b, bid)
                if kind == "h2" and first_h2:
                    h.paragraph_format.page_break_before = True  # 본문은 둘째 쪽부터
                    first_h2 = False
            elif kind == "p":
                if not lead_done:
                    lead_done = True
                    label(doc, "요약")
                    para(doc, b[1], size=S_LEAD, after=4, line=L_LEAD)
                    if heads:
                        contents(doc, heads)
                    continue
                p = para(doc, b[1])
                if b[1].startswith("**[추론]"):
                    tint(p, INFER_BOX)
            elif kind == "table":
                add_table(doc, b[1], body_width(doc))
            elif kind == "check":
                p = doc.add_paragraph()
                hang(p, Cm(0.6))
                p.paragraph_format.space_after, p.paragraph_format.line_spacing = Pt(3), Pt(L_CHECK)
                p._p.append(run_el("□\t", S_BODY, color=ACCENT))
                inline(p, b[1], S_BODY)
            elif kind == "bullet":
                p = doc.add_paragraph(style="List Bullet")
                p.paragraph_format.space_after = Pt(3)
                inline(p, b[1], S_BODY)
    n = doc.add_paragraph()
    n.paragraph_format.space_before = Pt(18)
    n._p.append(run_el(f"생성 문서 — 정본은 md({src_label})입니다. 고칠 때는 md 를 고친 뒤 "
                       f"depin/tools/build-report-docx.py 로 다시 만드세요.", 8, color=MUTE))
    return doc


if __name__ == "__main__":
    src = SRC.resolve()
    try:
        label_path = str(src.relative_to(REPO))
    except ValueError:
        label_path = src.name
    doc = build(src.read_text(encoding="utf-8"), SRC.stem, label_path)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"wrote {OUT}")
